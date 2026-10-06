# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)

def h(text, level=1):
    p = doc.add_heading(text, level=level)
    return p

def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = htext
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    return t

title = doc.add_heading('Weibull 分布研究成果发表策略与投稿指南', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
para('—— 基于贾祥《韦布尔分布及其可靠性统计方法》(科学出版社, 2021) 的后续论文规划').alignment = WD_ALIGN_PARAGRAPH.CENTER
para('编制日期：2026 年 10 月')

h('一、本书定位与研究基础', 1)
para('贾祥《韦布尔分布及其可靠性统计方法》（科学出版社，2021.3，约 32.8 万字）是本项目的方法底座。全书 7 章：')
table(['章', '内容', '方法深度'], [
    ['第1章', 'Weibull 分布及可靠性概述', '定义、寿命试验、可靠性指标'],
    ['第2章', '双参数 Weibull 的确定方法', '假设检验、拟合优度、分布误用分析、模型选择'],
    ['第3章', '基于极大似然估计的方法', '点估计、渐近正态 CI、枢轴量 CI、bootstrap CI'],
    ['第4章', '基于分布曲线拟合的方法', '失效概率估计、最小二乘/加权最小二乘、枢轴量 CI'],
    ['第5章', '基于 Bayes 的方法', '多源先验；MH-Gibbs 随机游走 MCMC（算法5.3）；R(t) 与剩余寿命 L'],
    ['第6章', '其他统计方法', '修正点估计、以分布函数为枢轴量的 CI、样本空间排序'],
    ['第7章', '改进 Weibull 分布', '7.1 q-Weibull（MLE+信息矩阵+渐近 CI）；7.2 指数化 Weibull（MLE+拟合）'],
])
para('关键判断：本书是"方法库"，但现代计算部分停留在 MLE+渐近正态 和 MH 随机游走 MCMC；未覆盖竞争风险、HMC/NUTS、弱可辨识性/正则化、后验预测检验、加速寿命外推、相依失效。这些空白即后续论文落点。', bold=True)

h('二、本书已覆盖 vs 当前研究空白', 1)
table(['方法维度', '本书已做', '本书未做（=可做论文）'], [
    ['q-Weibull 参数估计', 'MLE、Fisher/观测信息矩阵、渐近 CI', '贝叶斯/NUTS、竞争风险、弱可辨识性、正则化'],
    ['Weibull 贝叶斯计算', 'MH/Gibbs 随机游走 MCMC（混合差）', 'HMC/NUTS、ESS/divergence/R-hat、与 Stan 对比'],
    ['多源先验融合', '专家数据+历史寿命数据加权', 'NUTS 替代 MH 后的效率与区间校准'],
    ['竞争风险', '无', '独立 q-Weibull/Weibull、CIF、系统可靠度'],
    ['剩余寿命 RUL', '平均剩余寿命 Bayes 区间（MH）', 'NUTS 下 RUL 后验预测、PPC'],
    ['指数化 Weibull', '仅 MLE+曲线拟合', '贝叶斯 NUTS、与 Weibull/q-Weibull 模型选择'],
    ['加速寿命试验', '无', '步进/恒加 + Weibull/q-Weibull 外推'],
    ['相依失效', '独立潜在寿命假设', 'shared frailty / copula 敏感性'],
])

h('三、第 1 篇（已在投，系列锚点）', 1)
para('题目：q-Weibull Competing Risks under Weak Identifiability: Bayesian Inference via the No-U-Turn Sampler')
para('状态：LaTeX 源码 + 编译 PDF 已上传 Wiley 投稿系统，Cover Letter 已备；目标 QREI。')
para('核心故事：竞争截断削去尾部观测 → q 弱可辨识 → 先验收缩/正则化稳定点估计、贝叶斯真正增量在校准后验区间；尾部被观测或 q 远离先验中心时 MLE 更优。')

h('四、后续可做论文清单', 1)
para('第 2 篇：双参数 Weibull 多源先验融合的 NUTS 贝叶斯推断（最优先接着写）。', bold=True)
para('动机：本书第5章算法5.3（MH 随机游走）迹图自相关重、混合差。贡献：换 NUTS，在同一蓄电池/轴承数据上对比 ESS/s、divergence、R-hat、区间覆盖率；系统比较五类先验鲁棒性。目标：QREI/JSCS。')
para('第 3 篇：指数化 Weibull 的贝叶斯 NUTS 推断与模型选择。', bold=True)
para('动机：本书 7.2 仅 MLE。贡献：NUTS 推断指数化 Weibull，AIC/BIC+PPC 比较 Weibull/指数化 Weibull/q-Weibull。目标：Comm.Stat./JSCS。')
para('第 4 篇：竞争风险 RUL / 加速寿命外推（视前两篇意见）。', bold=True)
para('目标 IEEE TR/QREI。')
para('交叉方向：q-Weibull 重尾 → 保险损失/巨灾索赔建模，投 Insurance: Math. & Econ.。')

h('五、研究热点判断（2025—2026）', 1)
para('纯"构造新分布/堆参数"已过饱和；围绕既有灵活分布的方法论深化（贝叶斯计算效率、竞争风险、小样本多源先验融合、稳健推断、RUL）仍活跃。本书"MH 随机游走"旧基线用 NUTS 升级即最稳增量。')

h('六、目标期刊总览', 1)
table(['期刊', '定位', '匹配篇目'], [
    ['QREI (Wiley)', '工程可靠性，IF≈2.8，3区', '第1、2、4篇（首选）'],
    ['IEEE Trans. Reliability', '可靠性权威，双栏≤15页', '第4篇（冲刺）'],
    ['RESS', '顶刊，需真实案例', '第4篇'],
    ['JSCS (T&F)', '统计计算，≈1万词', '第2、3篇'],
    ['Comm. Stat.', '理论+模拟', '第3篇'],
    ['Insurance: Math. & Econ.', '重尾/保险风险', '交叉方向'],
    ['IJRQSE', 'IF≈1.4，审稿友好', '保底'],
    ['航空学报/数理统计与管理', 'EI/中文核心', '中文改写'],
])

h('七、必须正面区分的竞争论文', 1)
for s in [
    'q-Weibull 贝叶斯 Gibbs（Sigma J. Eng. Nat. Sci., 2026, 44(2):1094-1110）',
    '双模式竞争风险 HMC（Comm. Stat. T&M, 2026, 55(15):4860-4884）',
    'Weibull-Dhillon 竞争风险 HMC（QREI, 2025, 41(6):2624-2646）',
    'q-Weibull 应用综述（IEEE Trans. Reliab., 2025, 74(3)）',
    '本书第5章 MH-MCMC 算法5.3（第2篇旧基线）',
]:
    doc.add_paragraph(s, style='List Bullet')

h('八、投稿路线图', 1)
table(['顺序', '论文', '目标', '状态'], [
    ['第1篇', 'q-Weibull 竞争风险+弱可辨识性+NUTS', 'QREI', '已投'],
    ['第2篇', '双参数 Weibull 多源先验融合 NUTS', 'QREI/JSCS', '接着写'],
    ['第3篇', '指数化 Weibull 贝叶斯 NUTS+模型选择', 'Comm.Stat./JSCS', '第2篇投稿后'],
    ['第4篇', '竞争风险 RUL/加速寿命外推', 'IEEE TR/QREI', '视意见'],
])

h('九、投稿前检查清单', 1)
for s in [
    '同一套代码实机跑通 example 与 sim，结果可复现',
    '基线至少含 MLE、本书式 MH-MCMC（算法5.3）、NUTS',
    'NUTS 诊断齐全：R-hat、ESS/s、divergence 数与位置、max tree depth、BFMI',
    '覆盖率评估蒙特卡洛重复次数 R≥500',
    'bias/RMSE/coverage 覆盖全部参数，不只一个',
    '先验中心与先验类型做敏感性分析',
    '真实数据加 posterior predictive check，不夸大模型优势',
    '代码公开 GitHub，附 commit tag',
    '明确通讯作者 Huirong Cao（huirongcao@126.com）、单位、基金、利益冲突声明',
]:
    doc.add_paragraph(s, style='List Bullet')

doc.save(r'G:\OPT\Weibull发表策略\Weibull发表策略.docx')
print('OK')
